"""Build the final Chinese report from audited results, with no model inference.

Run with the bundled document Python after summarize_weekend_distillation.py
and plot_weekend_results.py have completed. Render and inspect before delivery.
"""
import argparse
import csv
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TITLE = 'TransMixer 周末融合与知识蒸馏实验全面报告'
CLASSES = ['健康草地','受胁草地','合成草地','树木','土壤','水体','住宅','商业区','道路','高速公路','铁路','停车场1','停车场2','网球场','跑道']


def main(out):
    data = json.loads((out/'周末实验总汇.json').read_text(encoding='utf-8'))
    assert data['status'] == 'completed'
    rows = data['legacy_methods']; by = {r['id']:r for r in rows}
    assert len(rows) == 10 and len(data['audit']['new_job_checks']) == 7
    agg = data['validation_protocol_21_jobs']['aggregates']
    delta = lambda a,b: by[a]['oa'] - by[b]['oa']
    f = lambda n: f'{n:.4f}'
    pp = lambda n: f'{n:+.4f}'
    judgement = lambda n: '提高' if n > 0 else '下降' if n < 0 else '持平'
    pages = []
    p = lambda text: ('p',text)
    h = lambda text: ('h',text)
    table = lambda headers,items,widths=None: ('table',headers,items,widths)
    fig = lambda name,caption: ('figure',name,caption)
    def page(title,blocks): pages.append((title,blocks))
    def metrics(ids):
        return [[k,by[k].get('method',by[k]['model']),f(by[k]['oa']),f(by[k]['aa']),f(by[k]['kappa']),f"{by[k]['parameters']:,}"] for k in ids]
    head = ['组','方法','OA %','AA %','Kappa ×100','参数']
    conclusion = (
        f"本轮 RSCNet 蒸馏让 F4 的 OA {judgement(delta('K2','F4'))} {abs(delta('K2','F4')):.4f} 个百分点；"
        f"相对无教师的标签平滑对照 {pp(delta('K2','C1'))} 个百分点。"
        + ('这组结果存在教师指导的正向信号，但仍需独立协议和更多种子确认。' if delta('K2','F4')>0 and delta('K2','C1')>0 else
           '当前不能据此认定教师指导优于普通正则化，所有正负结果都应保留。')
    )
    kd_ids=['K1','K2','K3','K4']
    if all(delta(k,'F0' if k=='K1' else 'F4')<0 for k in kd_ids):
        conclusion=(
            '四组蒸馏均未超过各自未蒸馏基线，本轮固定方案没有带来准确率提升。'
            f"F4 向 RSCNet 蒸馏后 OA 下降 {abs(delta('K2','F4')):.4f} 个百分点，"
            f"标签平滑相对 F4 为 {pp(delta('C1','F4'))} 个百分点。"
            '当前应保留轻量对照，暂不将普通蒸馏作为有效创新；这不等于否定所有可能的蒸馏方法。'
        )
    page('1 主要结论与实验边界',[
        p('本报告汇总 2026 年 10 月 24 日至 25 日保存的 TransMixer 实验，回答直接融合能否复现约 89% 的历史准确率、跨模态交互是否有效，以及知识蒸馏能否进一步提高轻量学生网络。正式训练共 31 次，按两套选模规则分开评价。'),
        p(conclusion),
        table(['问题','本轮证据','可以得出的结论'],[
            ['直接融合能否回到约 89%',f"F0 OA {f(by['F0']['oa'])}%；相对历史 88.9727 为 {pp(data['legacy_baseline_delta_vs_historical_pp'])} 个百分点",'接近历史水平，未要求逐位一致'],
            ['双向交互是否有效',f"F4−F0 {pp(delta('F4','F0'))}；F4−F1 {pp(delta('F4','F1'))} 个百分点",'本轮单种子描述性比较'],
            ['蒸馏是否改善直接融合',f"K1−F0 {pp(delta('K1','F0'))} 个百分点",judgement(delta('K1','F0'))],
            ['蒸馏是否改善交互模型',f"K2−F4 {pp(delta('K2','F4'))} 个百分点",judgement(delta('K2','F4'))],
            ['是否超出标签平滑效果',f"K2−C1 {pp(delta('K2','C1'))} 个百分点",'只能评价本次固定配置']],[4,9,4]),
        h('如何阅读这些分数'),
        p('OA 是全部测试中心的正确率；AA 是 15 类召回率的平均值；Kappa 衡量扣除随机一致后的表现，本报告统一乘以 100。准确率的相减使用百分点。新蒸馏对照只有 seed6，且沿用实验室用测试 OA 保存 checkpoint 的规则，因此属于探索结果，不是独立测试或统计显著性证明。'),
        p('本轮数据是 Houston2013 的 HSI 与 LiDAR。没有新增 Berlin 或 SAR 实验，不能将本报告的结论扩展到这些场景。')
    ])
    page('2 本周末做了哪些实验',[
        table(['批次','正式训练','数据与选模','用途'],[
            ['验证选模融合批次','F0 至 F6 × seed6 7 8 共 21 次','2267 训练 565 验证 12197 测试；验证 OA 选模','比较融合方向与随机性'],
            ['原实验室流程融合批次','F0 F1 F4 共 3 次','2832 官方训练 12197 测试；测试 OA 选模','复现历史直接融合口径'],
            ['新增教师与蒸馏批次','T1 T2 K1 K2 K3 K4 C1 共 7 次','沿用上述原流程；seed6；每次 100 轮','教师选择 蒸馏 置信过滤及平滑对照']],[4,5,5.5,2.5]),
        p('31 次只包含完成的正式训练。早期试启动 legacy_fusion_20261025_133300、合成输入检查、接口单元测试不计入训练总数。历史 9 月的单次成绩另列为背景，也不加入本周末统计。'),
        h('原流程统一设置'),
        table(['设置','本轮固定值'],[
            ['输入与预处理','PCA 30 维；11×11 窗口；全场景无标签 PCA 与标准化'],
            ['训练预算','100 轮；batch128；drop_last=True；每轮 22 个批次'],
            ['优化器','原 AdamW；学习率 0.0001；权重衰减 0.01'],
            ['随机种子与线程','seed6；OMP 与 MKL 线程均为 4'],
            ['模型保存','沿用实验室测试 OA 最优 checkpoint；不追加轮数或测试推理'],
            ['执行入口','原 scripts/train.py；原数据 训练和评估流程']],[4,13]),
        p('两套协议的训练中心数、选模依据和数值设置不同。不能将验证批次的 85% 与原流程的约 89% 直接相减，并把差额归因于网络或某一个设置。所有模块效果都应使用同一批次内的对照。')
    ])
    page('3 你的网络和教师分别在做什么',[
        h('直接融合与跨模态交互'),
        p('B0 的 HSI 分支先通过 3D 卷积和投影，再进入 TransMixer。块内按通道分出全局注意力路径与局部深度卷积路径；辅助 LiDAR 分支主要是 2D 卷积，并不是与 HSI 完全对称的局部全局编码器。最终将两种模态的特征按通道拼接，用 1×1 卷积、池化与分类头输出。'),
        p('F1 在辅助分支增加同类型自注意力，作为近似等参数容量对照。F4 在最终融合前加入双向交叉注意力，让 HSI 读取辅助信息，同时让辅助特征读取 HSI。两个方向都读取更新前特征，保留原融合头。F1 不严格匹配 F4 的全部计算量。'),
        h('教师和学生的关系'),
        p('教师是一张先独立训练好的分类网络。学生除了学习真实类别，还学习教师对 15 类给出的概率分布。例如教师把一个像素判断为道路 60%、高速公路 30%，它提供了这两类相似的信息。学生学习这种分布时仍保留真实标签损失，不是照抄教师的所有判断。'),
        table(['网络','训练时看到的输入','承担的角色'],[
            ['B0 与 F4 学生','PCA 后 HSI 加 LiDAR','最终用于推理的轻量网络'],
            ['RSCNet 教师 T1','原始 144 波段 HSI 加 PCA HSI 加 LiDAR','更丰富输入的训练指导'],
            ['CNN 教师 T2','PCA 后 HSI 加 LiDAR','与学生输入信息一致的教师对照']],[4,8,5]),
        p('教师的权重与 BatchNorm 状态在学生训练中冻结，并读取同一个增强后的 batch。学生均从头训练 100 轮，没有使用旧学生 checkpoint 继续训练。推理时仅保留原 B0 或 F4；教师不进入学生模型文件。'),
        p('RSCNet 与 CNN 的比较同时改变了网络结构和输入信息量。即使其中一位教师带来的收益更大，也不能只归因为该网络结构更适合蒸馏。')
    ])
    page('4 原版融合复现与历史结果',[
        table(head,metrics(['F0','F1','F4']),[.7,6,2.1,2.1,2.7,2.4]),
        table(['同批对照','OA 差值 百分点','解释'],[
            ['F1−F0',pp(delta('F1','F0')),'额外辅助变换带来的变化'],
            ['F4−F0',pp(delta('F4','F0')),'交互方案相对直接融合的总变化'],
            ['F4−F1',pp(delta('F4','F1')),'交互相对容量对照的变化']],[3,4,10]),
        p(f"F0 的 {f(by['F0']['oa'])}% 与历史 88.9727% 接近，相差 {pp(data['legacy_baseline_delta_vs_historical_pp'])} 个百分点。这回应了之前为何约为 89% 的问题：在恢复原 2832 个训练中心和原选模流程后，原模型确实回到了相近区间。但多个设置一起恢复，尚未逐项分离各设置的影响。"),
        h('历史单次结果仅作为背景'),
        table(['历史记录','OA %','证据边界'],[
            ['2026 年 9 月直接融合 B0','88.9727','历史会话保留的标准报告输出'],
            ['B1 局部残差','88.9399','历史单次特征提取消融'],
            ['B2 SSM 前残差','88.9317','历史单次特征提取消融'],
            ['B3 步长初始化','88.4644','历史单次特征提取消融'],
            ['B4 多尺度深度卷积','88.9891','相对旧 B0 约 +0.0164 个百分点']],[6,3,8]),
        p('这些旧记录不能并入本轮配对均值。B4 的极小单次差值并未证明稳定收益。关于历史 CNN 89.60% 与 RSCNet 90.01% 的记录，本轮已有重新训练的教师结果，下一页将以本轮实际值评价，避免把历史分数当作当前教师能力。')
    ])
    page('5 两位教师的实际表现',[
        table(head,metrics(['T1','T2','F0','F4']),[.7,6,2.1,2.1,2.7,2.4]),
        table(['教师对照','OA 差值 百分点','解读'],[
            ['T1 RSCNet−F4',pp(delta('T1','F4')),'主教师相对交互学生的独立分类表现'],
            ['T2 CNN−F4',pp(delta('T2','F4')),'同输入教师相对交互学生的表现'],
            ['T1−T2',pp(delta('T1','T2')),'含结构与输入信息量的共同差异']],[4,4,9]),
        p('两位教师在学生训练前分别完成 100 轮训练，并固定 checkpoint。教师选择基于实验室已有实现与历史记录；温度、权重和对照组在新结果出来前已经固定，没有依据学生测试成绩更换教师或搜索参数。'),
        p('教师独立准确率高，并不保证它传给学生的概率分布最容易学习；反过来，一位教师未超过学生，也不能提前保证蒸馏有益。是否值得使用，须看对应学生与不蒸馏基线的差值，并结合标签平滑对照。'),
        h('模型大小和部署代价'),
        p(f"本地 RSCNet 有 {by['T1']['parameters']:,} 个参数，CNN 有 {by['T2']['parameters']:,} 个。这里是本地实现的实际计数，不能替换成原论文参数数目。B0 学生为 {by['F0']['parameters']:,} 个参数，F4 为 {by['F4']['parameters']:,} 个。蒸馏后的学生与各自未蒸馏版本参数量相同。"),
        p('这种方法把额外成本放在训练阶段：先训练教师，再在学生每一步计算教师输出。本轮没有开展标准化推理计时，因此只能确认学生结构和参数量没有增加，不能进一步声称实际时延或能耗完全不变。')
    ])
    page('6 新增蒸馏对照如何设计',[
        table(['组','学生与训练方法','主要回答的问题'],[
            ['K1','B0 向 RSCNet 蒸馏','直接融合能否单独受益'],
            ['K2','F4 向 RSCNet 蒸馏','交互与蒸馏能否一起使用'],
            ['K3','F4 向 CNN 蒸馏','教师选择是否影响效果'],
            ['K4','F4 只学习高置信 RSCNet 输出','过滤不确定输出是否有帮助'],
            ['C1','F4 标签平滑 无教师','收益是否也能由简单正则化获得']],[1,8,8]),
        h('固定损失和参数'),
        p('普通组使用 L = CE + 0.3 × T² × KL，其中 T=4，KL 的方向是教师分布对学生分布。CE 让学生遵守真实标签；KL 让学生学习教师概率之间的关系；温度提高后，原本很小的类别概率更容易被学生观察到。'),
        p('K4 以温度 1 下教师最大概率是否达到 0.7 作为门槛，只在接受的样本上平均 KL。若一个 batch 全部不接受，则只使用 CE。这是置信过滤，不保证排除所有错误：教师也可能自信地预测错误。C1 使用 epsilon=0.1 的标签平滑，让真实标签保留较高权重，其他类别也获得少量目标概率。'),
        h('保留实验室流程'),
        p('训练器仅增加可选损失回调和结束诊断回调；默认训练仍使用原 CE。优化器仍只接收学生参数，数据增强、训练循环、测试、checkpoint 规则均沿用原流程。蒸馏系数、温度和阈值没有按测试分数继续调整。'),
        p('10 项接口测试与原生模型 GPU 预检已在正式蒸馏前通过，覆盖 KL 方向与温度缩放、零蒸馏权重、全拒绝、教师无梯度、配置兼容和保存恢复。它们验证实现能按设计运行，不代替性能实验。蒸馏公式参考 Hinton 等的原始工作，来源列于末页。')
    ])
    page('7 新增蒸馏实验的完整结果',[
        table(head,metrics(['F0','F4','K1','K2','K3','K4','C1']),[.7,6,2.1,2.1,2.7,2.4]),
        fig('关键对照差值.png','图 1  预先固定的 OA 对照差值  所有方法为同一 seed6 原流程'),
        p(conclusion),
        p(f"K3 相对 F4 为 {pp(delta('K3','F4'))} 个百分点。K4 相对普通 RSC 蒸馏 K2 为 {pp(delta('K4','K2'))} 个百分点。负值意味着该固定设置没有得到更高 OA，不能将失败组从报告中移除。AA 和 Kappa 同时列出，避免只看 OA 忽略类别不均衡。")
    ])
    diag = {k:by[k]['diagnostics'] for k in ['K1','K2','K3','K4','C1']}
    page('8 收益能归因于什么',[
        table(['关键比较','OA 差值 百分点','能够检验的范围'],[
            ['K1−F0',pp(delta('K1','F0')),'RSC 蒸馏对直接融合的效果'],
            ['K2−F4',pp(delta('K2','F4')),'RSC 蒸馏对交互模型的效果'],
            ['K2−K1',pp(delta('K2','K1')),'在相同教师指导下两种学生的差异'],
            ['K2−K3',pp(delta('K2','K3')),'两种教师的实用效果差异'],
            ['K4−K2',pp(delta('K4','K2')),'当前置信过滤规则的效果'],
            ['K2−C1',pp(delta('K2','C1')),'教师目标相对标签平滑的效果']],[3,4,10]),
        p('不能只凭“蒸馏后超过旧 B0”就认定蒸馏和交互都有效。K2 与 F4 控制学生结构，K1 与 F0 检查无交互时是否也受益；K2 与 K1 则反映有教师条件下，F4 相对 B0 的表现。它们共同帮助分离不同因素，但单次结果仍不能证明稳定的机制贡献。'),
        h('过滤实际启用了多少'),
        p(f"K4 在全部增强训练样本中的接受比例为 {diag['K4']['teacher_acceptance_fraction']*100:.2f}%，没有拒绝任何样本，因此本轮 0.7 阈值并未形成有效过滤。教师在这些训练 batch 上的正确率也是 100%，不能用它推断未见区域的可靠性。K4 与 K2 实际使用相同目标，二者 {pp(delta('K4','K2'))} 个百分点的差别不能解释为门控收益。旧流程没有采用严格确定性设置，具体数值差异来源未单独定位。"),
        table(['组','平均 CE','平均 KD','教师输出接受率'],[[k,f'{diag[k]["mean_ce"]:.5f}',f'{diag[k]["mean_kd"]:.5f}',f'{diag[k]["teacher_acceptance_fraction"]*100:.2f}%' if k!='C1' else '无教师'] for k in diag],[1.5,4,4,7.5]),
        p('标签平滑是实用的软目标正则化对照，不能严格匹配教师 KL 的所有梯度。本轮它优于四种蒸馏学生，但仍只增加了单次的小幅 OA，不能据此认定稳定有效。常规 KL 蒸馏和双向交互本身已有研究基础，不能仅因把两者组合就宣布论文创新成立。')
    ])
    change=[by['K2']['per_class_recall'][i]-by['F4']['per_class_recall'][i] for i in range(15)]
    order=sorted(range(15),key=lambda i:change[i],reverse=True)
    examples=order[:3]+order[-3:]
    page('9 哪些类别受益和受损',[
        fig('逐类召回率变化.png','图 2  各训练方法相对自身未蒸馏基线的逐类召回率变化'),
        table(['类别','F4 召回率','K2 召回率','K2−F4 百分点'],[[CLASSES[i],f(by['F4']['per_class_recall'][i]),f(by['K2']['per_class_recall'][i]),pp(change[i])] for i in examples],[5,4,4,4]),
        p('表中选取 K2−F4 排名前三与末三的类别，选择规则明确；完整 15 类指标保存在逐类召回率.csv。颜色表示召回率变化，不能解释为同一像素被纠正的比例，因为本报告复核的是混淆矩阵，没有重新运行推理或逐样本配对纠错。'),
        p('如果 OA 增加而某些类别下降，应如实保留。AA 有助于观察各类平均表现，但也不能替代具体业务类别要求。这里的类别分析是在看过测试结果后进行，只能为下一轮假设提供线索，不能反过来据此选择阈值并声称独立验证。')
    ])
    page('10 验证选模的三种子融合实验',[
        table(['组','方法','OA 均值±样本SD','AA 均值','Kappa 均值','相对 F0'],[[r['variant'],r['label'],f"{r['test_oa_mean']:.3f} ± {r['test_oa_sample_sd']:.3f}",f"{r['test_aa_mean']:.3f}",f"{r['test_kappa_mean']:.3f}",f"{r['paired_delta_mean']:+.3f}"] for r in agg],[.7,5,4,2.2,2.8,2.3]),
        fig('验证选模三种子.png','图 3  验证选模批次的配对种子结果  与本轮原流程分开解释'),
        p('该批 F4−F0 平均为 +0.440 个百分点，三个种子均为正；F4−F1 平均仅 +0.096，2 胜 1 负。F4 的 OA 标准差 3.173 大于 F0 的 2.883，不能说 F4 整体更稳定。F5 比交换分工 F6 平均低 0.265，当前实现没有支持原全局局部分工假设。'),
        p('565 个验证中心均落在某个训练窗口内，验证窗口也均与训练窗口相交；21 次中 20 次最佳验证 OA 达到 100%。窗口重叠本身不等于使用测试标签，但说明随机验证过易。原流程则直接用测试 OA 选模；两套规则各有限制，都没有完成空间隔离的独立泛化确认。')
    ])
    newrows=[by[k] for k in ['T1','T2','K1','K2','K3','K4','C1']]
    totalhours=sum(r['duration_seconds'] for r in newrows)/3600
    page('11 训练代价和科研定位',[
        table(['新增组','包含启动与预处理的作业耗时','模型推理参数'],[[r['id'],f"{r['duration_seconds']/60:.1f} 分钟",f"{r['parameters']:,}"] for r in newrows],[3,8,6]),
        p(f"新增 7 个作业总墙钟时间约 {totalhours:.2f} 小时，包括读取数据、PCA、每轮测试、保存模型和启动开销。它不是纯 GPU 计算时间，不宜与旧批次记录的训练时间直接相除。教师训练及教师前向均为新增成本。"),
        h('现在可以怎样汇报'),
        p(f"在 Houston2013 原实验室流程的 seed6 对照中，F4 相对直接融合提高 {pp(delta('F4','F0'))} 个百分点，相对容量对照提高 {pp(delta('F4','F1'))}；RSC 蒸馏 K2 相对 F4 为 {pp(delta('K2','F4'))}，相对标签平滑为 {pp(delta('K2','C1'))}。学生参数量保持不变。以上为探索结果，选模使用测试集，尚不能称为统计显著或 SOTA。"),
        h('后续研究优先级'),
        p('先确认最值得保留的固定方法能否在新的预设种子与空间分组验证下复现，再考虑跨数据集。若蒸馏与标签平滑持平，应优先保留更简单的方案；若置信过滤无收益，不继续堆叠相近门控。若教师明显更强而学生未改善，应先检查容量与软目标可学习性，不应直接更换大量超参数试到提升。'),
        p('本报告不自动追加这些训练。可发展的具体问题是轻量且不对称的模态分支如何避免传递不可靠信息，而非笼统地“加交互”或“加蒸馏”。已有分数只能帮助筛选假设，机制解释仍需要预先设计的可靠性对照。')
    ])
    page('12 数据来源与复核索引',[
        h('本地证据'),
        p('周末实验总汇.json 保存十组原流程指标、对照差值、21 次验证批次统计和核验结果。原流程十组结果.csv 可用于排序；逐类召回率.csv 保存完整类别比较；原始证据索引.csv 保存每组报告、checkpoint、配置及 SHA256。所有文件位于本报告所在目录。'),
        p('原三方法在加入蒸馏接口前已经封存，证据为原三方法最终验收.json。蒸馏批次的 manifest.json 保存接入后源码指纹；接入前原代码与蒸馏接口接入记录.json 保留修改记录。两批源码时期不同，不应因后续合法修改误判旧训练失真。'),
        p('新增七组均已核对原标准日志中的 100 轮完成记录、checkpoint 和 15×15 混淆矩阵。OA、AA、Kappa 由矩阵重新计算，每份测试总数均为 12197。五个学生对照均记录 2200 个损失步骤、281600 个增强样本；教师参数和 BN 状态不变、无梯度，且不进入学生 checkpoint。'),
        h('历史和论文来源'),
        p('历史 B0 与 B1 至 B4 背景来自本项目先前的综合评估报告及其证据目录，原路径 report/综合评估_20261025。历史 B0 的原始 checkpoint 尚未找回，不视为本轮重复训练。'),
        p('蒸馏原始论文：Hinton G 等，Distilling the Knowledge in a Neural Network，2015。https://arxiv.org/abs/1503.02531'),
        p('KL 实现接口：PyTorch functional.kl_div 官方文档。https://docs.pytorch.org/docs/stable/generated/torch.nn.functional.kl_div.html'),
        p('公开方法对照与既有交互文献详见先前综合评估报告。本轮 RSCNet 和 CNN 表格为本地实测；没有将论文作者分数混作本地 SOTA 排名。'),
        h('仍未获得的证据'),
        p('尚无蒸馏的多种子确认、空间隔离验证、独立测试、Berlin 或 SAR 泛化结果；没有标准化推理时延或能耗测量。整份报告包含成功与失败组，未因指标高低删除实验。Word 与 PDF 需经渲染逐页检查后才交付。')
    ])

    doc=Document();sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        s=doc.styles[name];s.font.name='Microsoft YaHei';s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',22),('Heading 1',16),('Heading 2',11),('Caption',8.5)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.bold=True;doc.styles['Subtitle'].font.italic=False
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
    for s in doc.styles:
        for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    sec.header.paragraphs[0].text='TransMixer 周末实验复盘  |  Houston2013';sec.header.paragraphs[0].style='Caption'
    foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('2026 年 10 月 25 日  ·  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for run in foot.runs:run.font.size=Pt(8)
    def add_table(headers,items,widths):
        widths=widths or [17/len(headers)]*len(headers);scale=17/sum(widths)
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Cm(w*scale)
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        for row in items:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for j,cell in enumerate(row.cells):
                cell.width=Cm(widths[j]*scale);tc=cell._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'60');x.set(qn('w:type'),'dxa');margins.append(x)
                tc.append(margins)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
                tc.append(borders)
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244A70');tc.append(sh)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing=1.1;paragraph.paragraph_format.space_after=Pt(2)
                    for run in paragraph.runs:
                        run.font.size=Pt(8.5);run.bold=i==0
                        if i==0:run.font.color.rgb=RGBColor(255,255,255)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    md=['# '+TITLE,'',f"生成依据：{data['created_at']} 的已核验汇总。",'']
    for i,(title,blocks) in enumerate(pages):
        if i:doc.add_page_break()
        else:
            doc.add_paragraph(TITLE,'Title')
            doc.add_paragraph('31 次正式训练  两套选模规则  全部正负结果','Subtitle')
        doc.add_paragraph(title,'Heading 1');md+=['## '+title,'']
        for block in blocks:
            if block[0] in ('p','h'):
                doc.add_paragraph(block[1],'Heading 2' if block[0]=='h' else 'Normal');md+=[('### ' if block[0]=='h' else '')+block[1],'']
            elif block[0]=='table':
                _,headers,items,widths=block;add_table(headers,items,widths)
                md+=['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']
                md+=['| '+' | '.join(map(str,row))+' |' for row in items];md+=['']
            else:
                _,name,caption=block
                doc.add_picture(str(out/name),width=Cm(16.8));doc.add_paragraph(caption,'Caption')
                md+=['!['+caption+']('+str(out/name).replace('\\','/')+')','']
    for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.core_properties.title=TITLE;doc.core_properties.author='科研实验复盘'
    doc.save(out/(TITLE+'.docx'));(out/(TITLE+'.md')).write_text('\n'.join(md),encoding='utf-8')
    with (out/'逐类召回率.csv').open('w',encoding='utf-8-sig',newline='') as file:
        w=csv.writer(file);w.writerow(['类别']+[r['id'] for r in rows])
        for i,name in enumerate(CLASSES):w.writerow([name]+[r['per_class_recall'][i] for r in rows])
    with (out/'原始证据索引.csv').open('w',encoding='utf-8-sig',newline='') as file:
        w=csv.writer(file);w.writerow(['组','原标准报告','checkpoint','配置','报告SHA256','checkpoint SHA256'])
        for r in rows:w.writerow([r['id'],r['report'],r['checkpoint'],r.get('config',''),r['report_sha256'],r.get('checkpoint_sha256','见原三方法最终验收.json')])
    (out/'报告生成记录.json').write_text(json.dumps({'title':TITLE,'planned_pages':len(pages),'formal_training_runs':31,'native_protocol_methods':10,'source_summary':str(out/'周末实验总汇.json'),'render_and_visual_review':'pending'},ensure_ascii=False,indent=2),encoding='utf-8')
    print(out/(TITLE+'.docx'))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--run',required=True,type=Path)
    main(parser.parse_args().run.resolve())
